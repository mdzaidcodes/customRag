from flask import Flask, render_template, request
from flask_socketio import SocketIO, emit
from sentence_transformers import SentenceTransformer, util
import pdfplumber
import os
import tempfile
import ollama
from docx import Document  # Import for reading .docx files
import traceback  # To help with detailed error logging

app = Flask(__name__)
socketio = SocketIO(app)

# Load the embedding model for creating embeddings
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

# Global placeholders for document data
document_chunks = []
document_embeddings = []

# Load the Llama model for response generation
model_name = 'llama2'
try:
    ollama.pull(model_name)
except Exception as e:
    print(f"Error preloading model '{model_name}': {e}")

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_document():
    # Check if the document file is in the request
    if 'document' not in request.files:
        return "No file uploaded", 400

    file = request.files['document']
    if file.filename == '':
        return "No file selected", 400

    # Save the uploaded file temporarily
    with tempfile.NamedTemporaryFile(delete=False) as temp_file:
        file.save(temp_file.name)
        temp_path = temp_file.name

    try:
        # Process the file based on its extension
        full_text = ""
        if file.filename.lower().endswith('.pdf'):
            # Extract text from a PDF
            with pdfplumber.open(temp_path) as pdf:
                full_text = " ".join(page.extract_text() for page in pdf.pages if page.extract_text())
        elif file.filename.lower().endswith('.docx'):
            # Extract text from a Word document
            doc = Document(temp_path)
            full_text = " ".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)
        else:
            return "Unsupported file format. Please upload a PDF or Word document.", 400

        if not full_text:
            return "Failed to extract text from the document. Please check the file content.", 400

        # Split the text into chunks and create embeddings
        chunk_size = 500
        chunks = [full_text[i:i + chunk_size] for i in range(0, len(full_text), chunk_size)]
        
        global document_chunks, document_embeddings
        document_chunks = chunks
        document_embeddings = embedding_model.encode(chunks, convert_to_tensor=True)

        return "Document uploaded and processed successfully", 200
    except Exception as e:
        print("Error processing document:", e)
        traceback.print_exc()  # Print the traceback for detailed debugging
        return f"Error processing document: {str(e)}", 500
    finally:
        os.remove(temp_path)  # Clean up the temporary file

@app.route('/delete', methods=['POST'])
def delete_document():
    global document_chunks, document_embeddings
    # Clear the document data from memory
    document_chunks = []
    document_embeddings = []
    return "Document deleted successfully", 200


@socketio.on('question')
def handle_question(data):
    question = data.get('text', '')

    # Check if there is an uploaded document
    if document_embeddings is None or document_embeddings.size(0) == 0:
        emit('answer', {'text': "No document uploaded or processed yet."})
        return

    # Encode the question for similarity matching
    question_embedding = embedding_model.encode(question, convert_to_tensor=True)

    # Find the most similar document chunk to the question
    try:
        similarities = util.pytorch_cos_sim(question_embedding, document_embeddings)
        best_match_idx = similarities.argmax().item()
        best_match_text = document_chunks[best_match_idx]

        # Prepare the prompt with the best-matching chunk as context
        prompt = (
            f"You are a helpful assistant. Answer questions strictly based on the following text. "
            f"Limit the output to 200 characters unless the user specifically asks it to be more than 200 characters. "
            f"If the question is outside this content, reply with 'I'm sorry, but I can only answer questions related to the provided text.' "
            f"Answer concisely: {best_match_text}"
        )

        # Generate response using the Llama model
        response = ollama.chat(
            model=model_name,
            messages=[
                {'role': 'system', 'content': prompt},
                {'role': 'user', 'content': question}
            ],
        )

        # Extract the generated response
        answer = response['message']['content']
        emit('answer', {'text': answer})
    except Exception as e:
        print("Error generating response:", e)
        traceback.print_exc()
        emit('answer', {'text': f'Error generating response: {str(e)}'})



if __name__ == '__main__':
    socketio.run(app, debug=True, use_reloader=False)
